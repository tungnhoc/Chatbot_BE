import io
import os
import json
import tempfile
import unittest
import subprocess
import sys
from pathlib import Path
from unittest.mock import patch
from sqlalchemy import create_engine, text, event
from sqlalchemy.orm import Session
from werkzeug.datastructures import FileStorage
from werkzeug.exceptions import NotFound, Conflict
from src.ingestion.readers import read_document
from src.ingestion.chunker import chunk_blocks
from src.evaluation.metrics import retrieval_metrics, citation_metrics
from src.evaluation.run import score


class ReadersTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_text_preserves_punctuation(self):
        path = self.root / 'test.txt'
        path.write_text('Giá: 125,5 tỷ đồng. Công thức: A+B.', encoding='utf-8')
        chunk = chunk_blocks(read_document(path, path.name))[0]
        self.assertIn('125,5', chunk['text'])
        self.assertIn('A+B', chunk['text'])

    def test_csv_headers(self):
        path = self.root / 'test.csv'
        path.write_text('Tên;Giá\nA;150000\nB;200000', encoding='utf-8')
        chunks = chunk_blocks(read_document(path, path.name))
        self.assertEqual(len(chunks), 2)
        self.assertTrue(all('Tên | Giá' in c['text'] for c in chunks))
        self.assertEqual(chunks[1]['metadata']['row'], 3)

    def test_docx_text_and_table(self):
        from docx import Document
        path = self.root / 'test.docx'
        doc = Document()
        doc.add_paragraph('Tiêu đề Việt Nam')
        table = doc.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, [('Tên', 'Giá'), ('A', '12,5')]):
            for cell, value in zip(row.cells, values):
                cell.text = value
        doc.save(path)
        blocks = read_document(path, path.name)
        self.assertEqual(blocks[0]['text'], 'Tiêu đề Việt Nam')
        self.assertEqual(blocks[1]['content_type'], 'table')

    def test_xlsx_sheet_and_header(self):
        from openpyxl import Workbook
        path = self.root / 'test.xlsx'
        book = Workbook()
        book.active.title = 'Gia'
        book.active.append(['Name', 'Price'])
        book.active.append(['A', 125.5])
        book.save(path)
        book.close()
        block = read_document(path, path.name)[0]
        self.assertEqual(block['source']['sheet'], 'Gia')
        self.assertIn('125.5', block['text'])

    def test_empty_and_unsupported(self):
        path = self.root / 'empty.txt'
        path.write_text('')
        with self.assertRaises(ValueError):
            read_document(path, path.name)
        with self.assertRaises(ValueError):
            read_document(path, 'image.png')

    def test_chunker_bounds(self):
        body = 'Một đoạn. ' * 300
        chunks = chunk_blocks([{'text': body, 'content_type': 'text', 'source': {'filename': 'a.md'}}])
        self.assertTrue(all(len(c['text']) <= 1000 for c in chunks))
        self.assertTrue(chunks[-1]['text'].endswith('đoạn.'))

    def pdf_fixture(self, name, content):
        # Minimal valid PDF built locally, with no external fixture/download dependency.
        stream = content.encode('ascii')
        objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
                   b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
                   b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
                   b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
                   b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
        data = b'%PDF-1.4\n'
        offsets = [0]
        for number, obj in enumerate(objects, 1):
            offsets.append(len(data))
            data += str(number).encode() + b' 0 obj\n' + obj + b'\nendobj\n'
        xref = len(data)
        data += b'xref\n0 6\n0000000000 65535 f \n'
        data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets[1:])
        data += f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF'.encode()
        path = self.root / name
        path.write_bytes(data)
        return path

    def test_pdf_text_and_table(self):
        path = self.pdf_fixture('table.pdf', '''BT /F1 12 Tf 50 740 Td (Report 125.5 USD) Tj ET
50 600 300 80 re S
50 640 m 350 640 l S
200 600 m 200 680 l S
BT /F1 12 Tf 60 655 Td (Name) Tj ET
BT /F1 12 Tf 210 655 Td (Price) Tj ET
BT /F1 12 Tf 60 615 Td (A) Tj ET
BT /F1 12 Tf 210 615 Td (125.5) Tj ET''')
        blocks = read_document(path, path.name)
        self.assertTrue(any('Report 125.5' in b['text'] for b in blocks))
        tables = [b for b in blocks if b['content_type'] == 'table']
        self.assertEqual(len(tables), 1)
        self.assertIn('Name | Price', tables[0]['text'])
        self.assertIn('125.5', tables[0]['text'])
        self.assertEqual(tables[0]['source']['page'], 1)

    def test_pdf_without_text_requires_ocr(self):
        path = self.pdf_fixture('blank.pdf', '50 50 100 100 re S')
        with self.assertRaisesRegex(ValueError, 'OCR'):
            read_document(path, path.name)


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.engine = create_engine('sqlite://')
        @event.listens_for(self.engine, 'connect')
        def configure(db, _):
            db.create_function('NOW', 0, lambda: '2026-10-19 12:00:00')
        self.session = Session(self.engine)
        statements = [
            'CREATE TABLE Conversations (ConversationID INTEGER PRIMARY KEY,UserID INTEGER,Title TEXT,CreatedAt DATETIME,UpdatedAt DATETIME)',
            'CREATE TABLE Documents (DocumentID INTEGER PRIMARY KEY,FileName TEXT,FilePath TEXT,FileType TEXT,UploadedBy INTEGER,FileSizeMB REAL,Description TEXT)',
            "CREATE TABLE IngestionJobs (JobID TEXT PRIMARY KEY,DocumentID INTEGER UNIQUE,UserID INTEGER,ContentHash TEXT,PipelineVersion TEXT,Status TEXT DEFAULT 'queued',Attempts INTEGER DEFAULT 0,Error TEXT,NextAttemptAt DATETIME DEFAULT CURRENT_TIMESTAMP, UNIQUE(UserID,ContentHash,PipelineVersion))",
            'CREATE TABLE ConversationDocuments (ConversationID INTEGER,DocumentID INTEGER,IsDeleted INTEGER DEFAULT 0,FileURL TEXT,UploadedAt DATETIME,PRIMARY KEY(ConversationID,DocumentID))',
            'CREATE TABLE DocumentChunks (ChunkID INTEGER PRIMARY KEY,DocumentID INTEGER,ChunkText TEXT,Metadata TEXT)',
            'CREATE TABLE Messages (MessageID INTEGER PRIMARY KEY,ConversationID INTEGER,Role TEXT,Text TEXT,Timestamp DATETIME DEFAULT CURRENT_TIMESTAMP)',
            'CREATE TABLE ChatTraces (TraceID TEXT,UserID INTEGER,ConversationID INTEGER,Payload TEXT,AssistantMessageID INTEGER UNIQUE)',
            "INSERT INTO Conversations VALUES (1,1,'Report',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)",
            "INSERT INTO Conversations VALUES (2,2,'Private',CURRENT_TIMESTAMP,CURRENT_TIMESTAMP)"]
        for sql in statements:
            self.session.execute(text(sql))
        self.session.commit()

    def tearDown(self):
        self.session.close()
        self.engine.dispose()
        self.tmp.cleanup()

    def upload(self, content=b'Price: 125.5 USD.'):
        from src.services.document_service import enqueue_document
        file = FileStorage(stream=io.BytesIO(content), filename='test.txt')
        with patch('src.services.document_service.storage_root', return_value=Path(self.tmp.name)):
            return enqueue_document(self.session, file, '1', 1)

    def test_upload_dedup(self):
        first, second = self.upload(), self.upload()
        self.assertEqual(first['document_id'], second['document_id'])
        self.assertEqual(self.session.execute(text('SELECT COUNT(*) FROM Documents')).scalar(), 1)
        self.assertEqual(len(list(Path(self.tmp.name).iterdir())), 1)

    def test_oversize_cleanup(self):
        with patch('src.services.document_service.MAX_FILE_BYTES', 5):
            with self.assertRaises(ValueError):
                self.upload()
        self.assertEqual(list(Path(self.tmp.name).iterdir()), [])

    def test_switch_to_gemini_creates_new_ingestion_job(self):
        with patch.dict(os.environ, {'EMBEDDING_BACKEND': 'huggingface'}):
            original = self.upload()
        with patch.dict(os.environ, {'EMBEDDING_BACKEND': 'gemini',
                'GEMINI_EMBEDDING_MODEL': 'gemini-embedding-2',
                'GEMINI_EMBEDDING_DIMENSION': '768', 'PINECONE_NAMESPACE': 'gemini-test'}):
            replacement = self.upload()
        self.assertNotEqual(original['document_id'], replacement['document_id'])
        self.assertEqual(self.session.execute(text('SELECT COUNT(*) FROM IngestionJobs')).scalar(), 2)

    def test_gemini_rejects_ready_documents_from_old_embedding_profile(self):
        from src.services.chat_service import handle_chat_query
        with patch.dict(os.environ, {'EMBEDDING_BACKEND': 'huggingface'}):
            original = self.upload()
        self.session.execute(text("UPDATE IngestionJobs SET Status='ready'"))
        self.session.commit()
        with patch.dict(os.environ, {'EMBEDDING_BACKEND': 'gemini',
                'GEMINI_EMBEDDING_MODEL': 'gemini-embedding-2',
                'GEMINI_EMBEDDING_DIMENSION': '768', 'PINECONE_NAMESPACE': 'gemini-test'}), \
                patch('src.services.chat_service.generate_answer') as llm:
            with self.assertRaises(Conflict):
                handle_chat_query(self.session, '1', 'Question', 1, [original['document_id']])
            llm.assert_not_called()

    def test_reupload_recovers_failed_document_after_original_was_lost(self):
        original = self.upload()
        source = self.session.execute(text('SELECT FilePath FROM Documents')).scalar_one()
        Path(source).unlink()
        self.session.execute(text("UPDATE IngestionJobs SET Status='failed',Attempts=3,Error='source_file_missing'"))
        self.session.commit()
        restored = self.upload()
        self.assertEqual(restored['document_id'], original['document_id'])
        self.assertEqual(restored['status'], 'queued')
        source = self.session.execute(text('SELECT FilePath FROM Documents')).scalar_one()
        self.assertEqual(Path(source).read_bytes(), b'Price: 125.5 USD.')
        row = self.session.execute(text('SELECT Attempts,Error FROM IngestionJobs')).one()
        self.assertEqual(tuple(row), (0, None))

    def test_worker_retry_stable_ids(self):
        from src.workers.ingestion import process_job
        result = self.upload()
        job = dict(self.session.execute(text(
            'SELECT * FROM IngestionJobs WHERE DocumentID=:doc'),
            {'doc': result['document_id']}).mappings().one())
        with patch('src.workers.ingestion.upsert_chunks', side_effect=RuntimeError('temporary')):
            with self.assertRaises(RuntimeError):
                process_job(self.session, job)
        first = self.session.execute(text('SELECT ChunkID FROM DocumentChunks')).scalars().all()
        with patch('src.workers.ingestion.upsert_chunks') as upsert:
            process_job(self.session, job)
            self.assertEqual(upsert.call_args.args[0], first)

    def test_chat_denies_foreign_conversation(self):
        from src.services.chat_service import handle_chat_query
        with patch('src.services.chat_service.generate_answer') as llm:
            with self.assertRaises(NotFound):
                handle_chat_query(self.session, '1', 'Question', 2)
            llm.assert_not_called()

    def test_chat_pending_rejected(self):
        from src.services.chat_service import handle_chat_query
        result = self.upload()
        with self.assertRaises(Conflict):
            handle_chat_query(self.session, '1', 'Question', 1, [result['document_id']])

    def test_chat_trace_sources(self):
        from src.services.chat_service import handle_chat_query
        result = self.upload()
        doc = result['document_id']
        self.session.execute(text("UPDATE IngestionJobs SET Status='ready'"))
        self.session.execute(text('INSERT INTO DocumentChunks VALUES (1,:doc,:body,:meta)'),
            {'doc': doc, 'body': 'Price: 125.5 USD.', 'meta': json.dumps({'filename': 'test.txt'})})
        self.session.commit()
        with patch('src.utils.pinecone_service.search_document_chunks', return_value=[{'chunk_id': 1, 'score': .9}]), patch('src.services.chat_service.generate_answer', return_value='125.5 USD [S1]'):
            reply = handle_chat_query(self.session, '1', 'Price?', 1, [doc])
            self.assertEqual(reply['sources'][0]['id'], 'S1')
        self.assertEqual(self.session.execute(text('SELECT COUNT(*) FROM Messages')).scalar(), 2)
        self.assertEqual(self.session.execute(text('SELECT COUNT(*) FROM ChatTraces')).scalar(), 1)
        from src.services.chat_service import load_conversation_history
        history = load_conversation_history(self.session, 1)
        self.assertEqual(history[-1]['sources'][0]['id'], 'S1')
        self.assertEqual(history[-1]['trace_id'], reply['trace_id'])

    def test_explicit_empty_document_selection(self):
        from src.services.chat_service import handle_chat_query
        self.upload()
        self.session.execute(text("UPDATE IngestionJobs SET Status='ready'"))
        self.session.commit()
        with patch('src.services.chat_service.generate_answer', return_value='No document selected'), patch('src.utils.pinecone_service.search_document_chunks') as search:
            result = handle_chat_query(self.session, '1', 'Question', 1, [])
            self.assertEqual(result['sources'], [])
            search.assert_not_called()

    def test_empty_ids_cannot_hide_unprocessed_attachment(self):
        from src.services.chat_service import handle_chat_query
        self.upload()
        with patch('src.services.chat_service.generate_answer') as llm:
            with self.assertRaises(Conflict):
                handle_chat_query(self.session, '1', 'Summarize this document', 1, [])
            llm.assert_not_called()

    def test_missing_job_does_not_silently_become_general_chat(self):
        from src.services.chat_service import handle_chat_query
        result = self.upload()
        self.session.execute(text('DELETE FROM IngestionJobs'))
        self.session.commit()
        with patch('src.services.chat_service.generate_answer') as llm:
            with self.assertRaises(Conflict):
                handle_chat_query(self.session, '1', 'Summarize this document', 1, [result['document_id']])
            llm.assert_not_called()

    def test_ready_document_without_retrieved_content_does_not_call_llm(self):
        from src.services.chat_service import handle_chat_query
        from werkzeug.exceptions import ServiceUnavailable
        result = self.upload()
        self.session.execute(text("UPDATE IngestionJobs SET Status='ready'"))
        self.session.commit()
        with patch('src.utils.pinecone_service.search_document_chunks', return_value=[]), patch('src.services.chat_service.generate_answer') as llm:
            with self.assertRaises(ServiceUnavailable):
                handle_chat_query(self.session, '1', 'Summarize this document', 1, [result['document_id']])
            llm.assert_not_called()

    def test_llm_failure_not_saved_as_success(self):
        from src.services.chat_service import handle_chat_query
        with patch('src.services.chat_service.generate_answer', side_effect=RuntimeError('unavailable')):
            with self.assertRaises(RuntimeError):
                handle_chat_query(self.session, '1', 'Question', 1)
        self.assertEqual(self.session.execute(text('SELECT COUNT(*) FROM Messages')).scalar(), 0)

    def api_client(self):
        from flask import Flask
        from flask_jwt_extended import JWTManager, create_access_token
        from src.routes.chat_routes import chat_bp
        from src.routes.upload_routes import upload_bp
        app = Flask(__name__)
        app.config['JWT_SECRET_KEY'] = 'test-secret-with-more-than-thirty-two-characters'
        JWTManager(app)
        app.register_blueprint(chat_bp)
        app.register_blueprint(upload_bp)
        with app.app_context():
            token = create_access_token(identity='1')
        return app.test_client(), {'Authorization': 'Bearer ' + token}

    def test_api_upload_202_and_foreign_denied(self):
        client, headers = self.api_client()
        from unittest.mock import MagicMock
        local = MagicMock(return_value=self.session)
        with patch('src.routes.upload_routes.SessionLocal', local), patch('src.services.document_service.storage_root', return_value=Path(self.tmp.name)):
            response = client.post('/api/conversations/1/documents', headers=headers,
                data={'files': (io.BytesIO(b'Hello'), 'a.txt')})
            self.assertEqual(response.status_code, 202)
            self.assertEqual(response.json['results'][0]['status'], 'queued')
            denied = client.post('/api/conversations/2/documents', headers=headers,
                data={'files': (io.BytesIO(b'Hello'), 'a.txt')})
            self.assertEqual(denied.status_code, 404)

    def test_api_foreign_files_and_trace_are_private(self):
        client, headers = self.api_client()
        self.session.execute(text("INSERT INTO Documents (DocumentID,FileName,UploadedBy) VALUES (99,'secret.txt',2)"))
        self.session.execute(text("INSERT INTO ConversationDocuments (ConversationID,DocumentID,FileURL) VALUES (2,99,'private-url')"))
        self.session.execute(text("INSERT INTO ChatTraces (TraceID,UserID,ConversationID,Payload) VALUES ('private',2,2,'{}')"))
        self.session.commit()
        with patch('src.routes.chat_routes.SessionLocal', return_value=self.session):
            response = client.get('/api/conversations/2/files', headers=headers)
            self.assertEqual(response.json['files'], [])
            trace = client.get('/api/chat/traces/private', headers=headers)
            self.assertEqual(trace.status_code, 404)


class MetricTests(unittest.TestCase):
    def test_numeric_facts_with_unicode_separators(self):
        from src.evaluation.metrics import numeric_fact_match
        self.assertEqual(numeric_fact_match('Price 112\u202f500 VND, 7 months [S1]', ['112500', '7']), 1)
        self.assertEqual(numeric_fact_match('P112500 costs 100 VND', ['112500']), 0)
    def test_retrieval(self):
        self.assertEqual(retrieval_metrics([{'page': 1}, {'page': 4}], [{'page': 1}, {'page': 2}]),
                         {'precision_at_k': .5, 'recall_at_k': .5})
        self.assertIsNone(retrieval_metrics([], [])['recall_at_k'])

    def test_citation_ids(self):
        self.assertEqual(citation_metrics('A [S1] B [S9]', [{'id': 'S1'}])['citation_id_validity'], .5)

    def test_missing_prediction(self):
        report = score([{'id': 'a', 'question': 'Q'}], [])
        self.assertEqual(report['overall']['error_rate'], 1)
        self.assertIsNone(report['overall']['correctness'])

    def test_mismatch(self):
        with self.assertRaises(ValueError):
            score([{'id': 'a', 'question': 'Q'}], [{'id': 'a', 'query': 'Wrong', 'answer': 'A'}])

    def test_cli_writes_json_and_csv(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / 'dataset.jsonl').write_text(json.dumps({'id': 'a', 'question': 'Q'}) + '\n', encoding='utf-8')
            (root / 'predictions.jsonl').write_text(json.dumps({'id': 'a', 'query': 'Q', 'answer': 'A', 'latency_ms': 20}) + '\n', encoding='utf-8')
            result = subprocess.run([sys.executable, '-m', 'src.evaluation.run',
                '--dataset', str(root / 'dataset.jsonl'), '--predictions', str(root / 'predictions.jsonl'),
                '--output', str(root / 'report')], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads((root / 'report.json').read_text(encoding='utf-8'))
            self.assertEqual(report['overall']['latency_p95_ms'], 20)
            self.assertIsNone(report['overall']['correctness'])
            self.assertTrue((root / 'report.csv').exists())


class LiveRegressionTests(unittest.TestCase):
    def test_llm_logging_works_with_ascii_stdout(self):
        from types import SimpleNamespace
        from src.llm.llm_client import generate_groq_response, get_last_usage
        response = SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content='Answer'))],
                                   usage=SimpleNamespace(model_dump=lambda: {'total_tokens': 42}))
        stream = io.TextIOWrapper(io.BytesIO(), encoding='ascii')
        with patch('src.llm.llm_client.Groq') as client, patch('sys.stdout', stream):
            client.return_value.chat.completions.create.return_value = response
            self.assertEqual(generate_groq_response('Q'), 'Answer')
            self.assertEqual(get_last_usage()['total_tokens'], 42)

    def test_empty_llm_output_is_failure(self):
        from types import SimpleNamespace
        from src.llm.llm_client import generate_groq_response, LanguageModelUnavailable
        with patch('src.llm.llm_client.Groq') as client:
            client.return_value.chat.completions.create.return_value = SimpleNamespace(
                choices=[SimpleNamespace(message=SimpleNamespace(content=''))], usage=None)
            with self.assertRaises(LanguageModelUnavailable):
                generate_groq_response('Q')

    def test_local_embeddings_mean_pool_ignores_padding(self):
        import numpy as np
        from types import SimpleNamespace
        from src.utils.local_embeddings import encode
        tokenizer = SimpleNamespace(encode_batch=lambda texts: [SimpleNamespace(ids=[1,2,0], attention_mask=[1,1,0], type_ids=[0,0,0])])
        model = SimpleNamespace(get_inputs=lambda: [SimpleNamespace(name='input_ids'), SimpleNamespace(name='attention_mask')],
            run=lambda *args: [np.stack([np.ones(384), np.ones(384)*3, np.ones(384)*100])[None, :, :]])
        with patch('src.utils.local_embeddings.resources', return_value=(tokenizer, model)):
            vectors = encode(['hello'])
        self.assertEqual(vectors.shape, (1,384))
        self.assertTrue(np.all(vectors == 2))

    def test_local_backend_requires_separate_namespace(self):
        from src.utils.pinecone_service import _namespace
        with patch.dict('os.environ', {'EMBEDDING_BACKEND': 'local-onnx', 'PINECONE_NAMESPACE': ''}):
            with self.assertRaises(ValueError):
                _namespace()

    def test_synthetic_corpus_is_repeatable(self):
        from src.evaluation.fixtures import generate
        import hashlib
        with tempfile.TemporaryDirectory() as folder:
            files, cases = generate(folder)
            before = [hashlib.sha256((Path(folder)/name).read_bytes()).hexdigest() for name in files]
            generate(folder)
            after = [hashlib.sha256((Path(folder)/name).read_bytes()).hexdigest() for name in files]
            self.assertEqual(before, after)
            self.assertEqual(len(cases), 52)


if __name__ == '__main__':
    unittest.main()
