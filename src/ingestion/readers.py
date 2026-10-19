"""Read supported documents without discarding punctuation or table headers."""
import csv
import io
import zipfile
from itertools import islice
from pathlib import Path

SUPPORTED = {'.pdf', '.docx', '.txt', '.md', '.csv', '.xlsx'}
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_BLOCKS = 20000


def _table(rows):
    rows = [[str(v).replace('\n', ' ').replace('|', '\\|') if v is not None else '' for v in r] for r in rows]
    if not rows:
        return []
    header = ' | '.join(rows[0])
    # One row per block: every chunk retains column names.
    return [header + '\n' + ' | '.join(row) for row in rows[1:] if any(row)] or [header]


def read_document(path, filename):
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError('Unsupported document format')
    blocks = []

    def add(value, kind='text', **source):
        if value and value.strip():
            blocks.append({'text': value.strip(), 'content_type': kind,
                           'source': {'filename': filename, **source}})
            if len(blocks) > MAX_BLOCKS:
                raise ValueError('Document contains too many blocks')

    if ext in {'.docx', '.xlsx'}:
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError('Expanded document exceeds 100 MiB')
    if ext in {'.txt', '.md'}:
        add(Path(path).read_text(encoding='utf-8-sig'))
    elif ext == '.csv':
        with open(path, encoding='utf-8-sig', newline='') as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=',;\t')
            except csv.Error:
                dialect = csv.excel
            rows = list(islice(csv.reader(handle, dialect), MAX_BLOCKS + 2))
            if len(rows) > MAX_BLOCKS + 1:
                raise ValueError('CSV contains too many rows')
            for number, row in enumerate(_table(rows), 2):
                add(row, 'table', row=number)
    elif ext == '.docx':
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        table_no = 0
        for element in document.element.body:
            if element.tag.endswith('}p'):
                add(Paragraph(element, document).text)
            elif element.tag.endswith('}tbl'):
                table_no += 1
                for number, row in enumerate(_table([[c.text for c in r.cells] for r in Table(element, document).rows]), 2):
                    add(row, 'table', table=table_no, row=number)
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                if (sheet.max_row or 0) > MAX_BLOCKS + 1 or (sheet.max_column or 0) > 1000:
                    raise ValueError('Spreadsheet dimensions exceed supported limits')
                rows = sheet.iter_rows(values_only=True)
                header = next(rows, None)
                if header is None:
                    continue
                for number, row in enumerate(rows, 2):
                    if any(v is not None for v in row):
                        add(_table([header, row])[0], 'table', sheet=sheet.title, row=number)
        finally:
            workbook.close()
    elif ext == '.pdf':
        import pdfplumber
        with pdfplumber.open(path) as document:
            if len(document.pages) > 500:
                raise ValueError('PDF exceeds 500 pages')
            for number, page in enumerate(document.pages, 1):
                tables = page.find_tables()
                # Remove table characters from prose to avoid duplicate/conflicting context.
                def outside_tables(obj):
                    x = (obj.get('x0', 0) + obj.get('x1', 0)) / 2
                    y = (obj.get('top', 0) + obj.get('bottom', 0)) / 2
                    return not any(t.bbox[0] <= x <= t.bbox[2] and t.bbox[1] <= y <= t.bbox[3] for t in tables)
                add(page.filter(outside_tables).extract_text() or '', page=number)
                for index, table in enumerate(tables, 1):
                    for row_no, row in enumerate(_table(table.extract()), 2):
                        add(row, 'table', page=number, table=index, row=row_no)
    if not blocks:
        raise ValueError('No extractable text. Scans/images require OCR, which is not enabled.')
    return blocks
