"""Deterministic synthetic corpus; not a domain-quality benchmark."""
import csv
import json
from pathlib import Path


def write_pdf(path, lines):
    content = '\n'.join(f'BT /F1 12 Tf 50 {740-i*22} Td ({line}) Tj ET' for i, line in enumerate(lines)).encode('ascii')
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length '+str(len(content)).encode()+b' >>\nstream\n'+content+b'\nendstream']
    data, offsets = b'%PDF-1.4\n', []
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data += f'{index} 0 obj\n'.encode()+obj+b'\nendobj\n'
    start = len(data)
    data += b'xref\n0 6\n0000000000 65535 f \n'
    data += b''.join(f'{offset:010d} 00000 n \n'.encode() for offset in offsets)
    data += f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()
    path.write_bytes(data)


def generate(root):
    from docx import Document
    from openpyxl import Workbook
    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    previous = root / 'dataset.jsonl'
    if previous.exists():
        cases = [json.loads(line) for line in previous.read_text(encoding='utf-8').splitlines() if line.strip()]
        files = list(dict.fromkeys(name for case in cases for name in case['filenames']))
        if all((root/name).is_file() for name in files):
            return files, cases
    cases, files, products = [], [], []
    for group, extension in enumerate(['pdf', 'docx', 'txt', 'md', 'csv', 'xlsx'], 1):
        filename = f'catalog_{group}.{extension}'
        path = root / filename
        rows, lines = [], []
        for index in range(1, 5):
            code = f'P{group}{index}'
            price, warranty = 100000 * group + 12500 * index, 6 * group + index
            rows.append([code, price, warranty])
            lines.append(f'Product {code}: price {price} VND; warranty {warranty} months.')
            source = {'filename': filename}
            if extension == 'pdf':
                source['page'] = 1
            elif extension in {'csv', 'xlsx', 'docx'}:
                source['row'] = index + 1
            if extension == 'xlsx':
                source['sheet'] = 'Prices'
            products.append((filename, code, price, warranty, source))
            for field, expected, unit in [('giá', price, 'VND'), ('thời hạn bảo hành', warranty, 'tháng')]:
                cases.append({'id': f'{extension}-{index}-{unit}', 'category': extension,
                    'question': f'Theo tài liệu, {field} của sản phẩm {code} là bao nhiêu?',
                    'reference_answer': f'{expected} {unit}.', 'expected_values': [str(expected)],
                    'relevant_sources': [source], 'filenames': [filename]})
        headers = ['Product', 'Price VND', 'Warranty months']
        if extension == 'pdf':
            write_pdf(path, lines)
        elif extension in {'txt', 'md'}:
            path.write_text('\n\n'.join(lines), encoding='utf-8')
        elif extension == 'csv':
            with path.open('w', newline='', encoding='utf-8') as handle:
                writer = csv.writer(handle)
                writer.writerow(headers)
                writer.writerows(rows)
        elif extension == 'docx':
            document = Document()
            document.add_heading('Product catalog', 1)
            table = document.add_table(rows=1, cols=3)
            for cell, value in zip(table.rows[0].cells, headers):
                cell.text = value
            for row in rows:
                for cell, value in zip(table.add_row().cells, row):
                    cell.text = str(value)
            document.save(path)
        else:
            workbook = Workbook()
            workbook.active.title = 'Prices'
            workbook.active.append(headers)
            for row in rows:
                workbook.active.append(row)
            workbook.save(path)
            workbook.close()
        files.append(filename)
    for index, pair in enumerate([(products[0], products[-1]), (products[4], products[8])], 1):
        a, b = pair
        cases.append({'id': f'multi-{index}', 'category': 'multi_document',
            'question': f'Giá của {a[1]} và {b[1]} lần lượt là bao nhiêu VND?',
            'reference_answer': f'{a[1]}: {a[2]} VND; {b[1]}: {b[2]} VND.',
            'expected_values': [str(a[2]), str(b[2])], 'relevant_sources': [a[4], b[4]],
            'filenames': [a[0], b[0]]})
    for index, question in enumerate(['Doanh thu năm 2030 là bao nhiêu?', 'Địa chỉ nhà máy sản xuất là gì?'], 1):
        cases.append({'id': f'missing-{index}', 'category': 'unanswerable', 'question': question,
            'reference_answer': 'Tài liệu không cung cấp thông tin này.', 'relevant_sources': [], 'filenames': files})
    dataset = root / 'dataset.jsonl'
    dataset.write_text(''.join(json.dumps(c, ensure_ascii=False)+'\n' for c in cases), encoding='utf-8')
    return files, cases
